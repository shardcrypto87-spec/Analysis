"""
Batch-parse every newly supplied WTB file, categorize via the GIFI master
mapping (+ entity-scoped Map No fallback), and cross-validate FY2025 net
income against the independently-validated Validation Log tab in
GIFI_Master_Mapping_v2.xlsx. Entities that don't match within rounding are
flagged for manual investigation rather than silently accepted (same
discipline used for Restaurant IndiaRosa 2).
"""
import json
import openpyxl
from collections import defaultdict

from entity_map import ENTITIES, FILE_PREFIX_TO_ENTITY, UNEXPECTED_ENTITIES

GIFI_XLSX = "data/GIFI_Master_Mapping_v2.xlsx"


def norm_mapno(s):
    """Collapse all whitespace so '40.1' and '40. 1' match the same group
    (the source WTB is inconsistent about the space after the period)."""
    return "".join(str(s).split()) if s is not None else ""


def load_gifi_map():
    wb = openpyxl.load_workbook(GIFI_XLSX, data_only=True)
    gifi_map = {}
    for code, cat, _ in wb["GIFI Master Mapping"].iter_rows(min_row=2, values_only=True):
        if code is not None:
            gifi_map[str(code).strip()] = cat

    fallback = {}
    started = False
    for row in wb["Fallback - Map No (blank GIFI)"].iter_rows(values_only=True):
        if row[0] == "FilePrefix":
            started = True
            continue
        if started and row[0]:
            fp, mapno, cat, _ = row
            fallback[(fp, norm_mapno(mapno))] = cat

    validation_log = {}
    for row in wb["Validation Log"].iter_rows(min_row=2, values_only=True):
        if row[0]:
            validation_log[row[0]] = row[1]

    return gifi_map, fallback, validation_log


def categorize(gifi, mapno, file_prefix, gifi_map, fallback):
    gifi_str = str(gifi).strip() if gifi is not None else ""
    if gifi_str:
        cat = gifi_map.get(gifi_str)
        if cat:
            return cat
    mapno_str = norm_mapno(mapno)
    return fallback.get((file_prefix, mapno_str), "UNCATEGORIZED")


def parse_wtb(path, file_prefix, gifi_map, fallback):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True))
    inc = [r for r in rows if len(r) > 4 and r[4] == "Income statement"]

    totals = {"FY2025": defaultdict(float), "FY2024": defaultdict(float), "FY2023": defaultdict(float)}
    uncategorized = []
    for r in inc:
        acct, name, mapno, gifi = r[0], r[1], r[3], r[13]
        final_cy, final_py, final_py2 = r[12], r[14], r[15]
        cat = categorize(gifi, mapno, file_prefix, gifi_map, fallback)
        if cat == "UNCATEGORIZED" and (final_cy or final_py or final_py2):
            uncategorized.append((acct, name, final_cy, final_py, final_py2))
        totals["FY2025"][cat] += final_cy or 0
        totals["FY2024"][cat] += final_py or 0
        totals["FY2023"][cat] += final_py2 or 0

    net_income = {fy: round(-sum(v for k, v in totals[fy].items() if k != "UNCATEGORIZED"), 2)
                  for fy in totals}
    return totals, net_income, uncategorized


def main():
    gifi_map, fallback, validation_log = load_gifi_map()

    import glob
    import os
    wtb_files = glob.glob("data/*_WTB.xlsx")
    file_prefixes_present = set()
    for path in wtb_files:
        base = os.path.basename(path)
        prefix = base.replace("_2025_WTB.xlsx", "")
        file_prefixes_present.add(prefix)

    results = {}
    for prefix in sorted(file_prefixes_present):
        if prefix == "Res_In2":
            continue  # already fully rebuilt with account-level exceptions
        entity_id = FILE_PREFIX_TO_ENTITY.get(prefix)
        path = f"data/{prefix}_2025_WTB.xlsx"
        if entity_id is None:
            print(f"!! {prefix}: no entity mapping (unexpected file), skipping — {path}")
            continue
        totals, net_income, uncategorized = parse_wtb(path, prefix, gifi_map, fallback)
        entity_name = ENTITIES[entity_id][0]
        # Validation Log keys are like "9357-7427 QI" or "Bistro Guru Inc.";
        # "Sandhu & Sandhu Enr." is logged there as just "Sandhu & Sandhu".
        bench = validation_log.get(entity_id) or validation_log.get(entity_name)
        if bench is None and entity_id == "N/A_SanSan":
            bench = validation_log.get("Sandhu & Sandhu")
        computed = net_income["FY2025"]
        match = None
        if bench is not None:
            match = abs(computed - bench) < 1.0
        results[entity_id] = {
            "prefix": prefix,
            "entity_name": entity_name,
            "totals": {fy: dict(totals[fy]) for fy in totals},
            "net_income": net_income,
            "uncategorized": uncategorized,
            "validation_benchmark": bench,
            "validation_match": match,
        }
        status = "MATCH" if match else ("MISMATCH" if match is False else "no benchmark")
        print(f"{prefix:10s} -> {entity_name:30s} FY2025 NI={computed:>14,.2f}  benchmark={bench}  [{status}]")
        if uncategorized:
            print(f"   uncategorized w/ balance: {uncategorized}")

    with open("output/all_entities_extract.json", "w") as f:
        json.dump(results, f, indent=2, default=str)

    print()
    print("Unexpected files (not in handoff's 18-entity list):", UNEXPECTED_ENTITIES)


if __name__ == "__main__":
    main()
