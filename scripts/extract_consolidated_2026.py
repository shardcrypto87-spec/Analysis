"""
Parse the consolidated CaseWare export (Sandhu_G__2026.xlsx) — one workbook,
16 tabs, one per entity, same column layout as the individual WTB exports.

Like every individual WTB file seen in this project, the column HEADER
dates are stale by one year (confirmed by cross-checking against the
already-validated FY2025 benchmarks and against the standalone
Res_Ind_2026_WTB.xlsx file): the column printed "Final: 2025-03-31" is
actually FY2026, "Prior: 2024-03-31" is actually FY2025, etc. Column
POSITION is trusted, never the header text, exactly as elsewhere in this
project.

Produces FY2026 fact rows for 15 of the 16 entities.

Sandhu & Sandhu Enr. ("Sheet16" — never renamed by CaseWare, unlike the
other 15 tabs) is EXCLUDED: its Final/Prior/Prior2/Prior3 columns are a
byte-for-byte duplicate of the old standalone San_San_2025_WTB.xlsx file
(confirmed account-by-account), not refreshed data. So its "Final" column
here is last year's FY2025 again, not FY2026 — there is no genuine FY2026
data for this entity in this file. The FY2025 figure already in the
workbook (-$15,821.17, from the standalone file, matching the Validation
Log) stays as-is; an earlier read of this file's "Prior" column
(-$11,141.31) as a restated FY2025 was wrong — it's this same stale
tab's FY2024, not a new FY2025. A fresh San_San export is needed before
FY2026 (or any correction to FY2025) can be added for this entity.
"""
import json
from collections import defaultdict

import openpyxl

from entity_map import FILE_PREFIX_TO_ENTITY

GIFI_XLSX = "data/GIFI_Master_Mapping_v2.xlsx"
SRC_XLSX = "data/Sandhu_G__2026.xlsx"

# Sheet name in the consolidated file -> canonical FilePrefix (as used by
# FILE_PREFIX_TO_ENTITY and the GIFI fallback table). Two sheets don't match
# the canonical spelling directly: "Lob_lou" (lowercase l) and "Sheet16",
# which CaseWare/Excel never renamed — confirmed as Sandhu & Sandhu Enr. by
# its distinctive account names (individual capital accounts, "Loan payable
# - Ajmer Sandhu Family Trust"), not present in any other entity.
SHEET_TO_PREFIX = {
    "Res_Ind": "Res_Ind", "Res_In2": "Res_In2", "Res_In3": "Res_In3",
    "Res_Her": "Res_Her", "Res_Sa2": "Res_Sa2", "Aub_Stl": "Aub_Stl",
    "Lob_lou": "Lob_Lou", "Bis_Gur": "Bis_Gur", "936_167": "936_167",
    "942_454": "942_454", "935_742": "935_742", "945_420": "945_420",
    "945_417": "945_417", "945_416": "945_416", "San_Lea": "San_Lea",
    "Sheet16": "San_San",
}

# Same 6 account-level overrides used for Res_In2 in extract_indiarosa2.py —
# still apply; nothing suggests CaseWare fixed the underlying miscoding.
IN2_EXCEPTIONS = {
    45130: "Marketing & Entertainment",
    45125: "Marketing & Entertainment",
    44215: "Other Operating",
    33455: "Other Operating",
    45250: "Other Operating",
    45150: "Other Operating",
}


def norm_mapno(s):
    return "".join(str(s).split()) if s is not None else ""


def load_gifi_map():
    wb = openpyxl.load_workbook(GIFI_XLSX, data_only=True)
    gifi_map = {}
    for code, cat, _ in wb["GIFI Master Mapping"].iter_rows(min_row=2, values_only=True):
        if code is not None:
            gifi_map[str(code).strip()] = cat

    fallback = {}
    started = False
    for row in wb["Fallback - Map No (blank GIFI)"].iter_rows(values_only=True):
        if row[0] == "FilePrefix":
            started = True
            continue
        if started and row[0]:
            fp, mapno, cat, _ = row
            fallback[(fp, norm_mapno(mapno))] = cat
    return gifi_map, fallback


def categorize(acct, gifi, mapno, file_prefix, gifi_map, fallback):
    if file_prefix == "Res_In2" and acct in IN2_EXCEPTIONS:
        return IN2_EXCEPTIONS[acct]
    gifi_str = str(gifi).strip() if gifi is not None else ""
    if gifi_str:
        cat = gifi_map.get(gifi_str)
        if cat:
            return cat
    return fallback.get((file_prefix, norm_mapno(mapno)), "UNCATEGORIZED")


def main():
    gifi_map, fallback = load_gifi_map()
    wb = openpyxl.load_workbook(SRC_XLSX, data_only=True)

    fy2026_rows = []
    uncategorized = []
    net_income = {}
    # Full 4-year re-derivation per entity, for the cross-check against the
    # existing final_fact_table.json (FY2023-FY2025) below.
    recheck_totals = {}

    for sheet_name, prefix in SHEET_TO_PREFIX.items():
        entity_id = FILE_PREFIX_TO_ENTITY[prefix]
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True))
        inc = [r for r in rows if len(r) > 4 and r[4] == "Income statement"]

        # idx: 0 AccountNo, 1 Name, 3 MapNo, 4 Type, 12 Final(=FY2026),
        # 13 GIFI, 14 Prior(=FY2025), 15 Prior2(=FY2024), 16 Prior3(=FY2023).
        totals = {"FY2026": defaultdict(float), "FY2025": defaultdict(float),
                  "FY2024": defaultdict(float), "FY2023": defaultdict(float)}
        for r in inc:
            acct, name, mapno, gifi = r[0], r[1], r[3], r[13]
            f26, f25, f24, f23 = r[12], r[14], r[15], r[16]
            cat = categorize(acct, gifi, mapno, prefix, gifi_map, fallback)
            if cat == "UNCATEGORIZED" and (f26 or f25 or f24 or f23):
                uncategorized.append((sheet_name, prefix, acct, name, f26, f25, f24, f23))
            totals["FY2026"][cat] += f26 or 0
            totals["FY2025"][cat] += f25 or 0
            totals["FY2024"][cat] += f24 or 0
            totals["FY2023"][cat] += f23 or 0

        recheck_totals[entity_id] = totals

        if entity_id == "N/A_SanSan":
            continue  # stale duplicate tab — no genuine FY2026 data, see module docstring

        ni26 = round(-sum(v for k, v in totals["FY2026"].items() if k != "UNCATEGORIZED"), 2)
        net_income[entity_id] = ni26

        for cat, amt in totals["FY2026"].items():
            if cat == "UNCATEGORIZED" or abs(amt) < 0.005:
                continue
            fy2026_rows.append([entity_id, cat, "FY2026", round(amt, 2)])

    print("FY2026 net income by entity (new data, not previously in the workbook):")
    for eid, ni in net_income.items():
        print(f"  {eid:15s} {ni:>14,.2f}")

    print()
    print("Uncategorized rows with a nonzero balance (should be empty):", uncategorized)

    # ---- Cross-check: FY2023-FY2025 re-derived here vs. the existing,
    # already-validated final_fact_table.json — must match to the cent for
    # every entity (San_San excluded — its tab is a stale duplicate, not
    # new data, see module docstring) before this file is trusted. ----
    with open("output/final_fact_table.json") as f:
        existing = json.load(f)
    existing_ni = defaultdict(lambda: defaultdict(float))
    for eid, cat, fy, amt in existing["fact_rows"]:
        existing_ni[eid][fy] += amt

    print()
    print("Cross-check vs existing final_fact_table.json (FY2023-FY2025):")
    all_match = True
    for entity_id, totals in recheck_totals.items():
        if entity_id == "N/A_SanSan":
            continue  # stale duplicate tab, not a meaningful cross-check
        for fy in ("FY2023", "FY2024", "FY2025"):
            new_ni = round(-sum(v for k, v in totals[fy].items() if k != "UNCATEGORIZED"), 2)
            old_ni = round(-existing_ni[entity_id][fy], 2)
            match = abs(new_ni - old_ni) < 0.01
            if not match:
                all_match = False
                print(f"  MISMATCH {entity_id:15s} {fy}  new={new_ni:>14,.2f}  old={old_ni:>14,.2f}")
    print("  All 15 entities match exactly (San_San excluded, stale tab)." if all_match else "  See mismatches above.")

    out = {
        "fy2026_rows": fy2026_rows,
        "net_income_fy2026": net_income,
        "uncategorized": uncategorized,
        "excluded_entities": {"N/A_SanSan": "Sheet16 is a byte-for-byte duplicate of the old "
                               "standalone San_San_2025_WTB.xlsx — no genuine FY2026 data present"},
    }
    with open("output/consolidated_2026_extract.json", "w") as f:
        json.dump(out, f, indent=2, default=str)
    print()
    print(f"Wrote output/consolidated_2026_extract.json — {len(fy2026_rows)} FY2026 rows "
          f"(15 entities; San_San excluded).")


if __name__ == "__main__":
    main()
